"""
NetGuard-IDS Presentation Deck Generator.
Generates an executive, dark-themed, 10-slide PowerPoint (.pptx) presentation
perfect for presenting in front of professors and evaluators.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DOCS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
ARCH_IMG = os.path.join(DOCS_DIR, "architecture_diagram.png")
OUTPUT_PPTX = os.path.join(DOCS_DIR, "Presentation.pptx")

# Cyber Theme Palette
COLOR_BG = RGBColor(14, 17, 23)        # #0E1117
COLOR_CARD = RGBColor(26, 29, 36)      # #1A1D24
COLOR_TEAL = RGBColor(0, 212, 178)     # #00D4B2
COLOR_WHITE = RGBColor(245, 245, 245)
COLOR_MUTED = RGBColor(160, 165, 175)
COLOR_CRIT = RGBColor(255, 77, 77)

def apply_background(slide):
    """Fill slide with dark modern background."""
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title, category="NETGUARD-IDS DEFENSIVE FRAMEWORK"):
    """Consistent header style across slides."""
    # Category tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_TEAL

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(8.4), Inches(0.7))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_WHITE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==================== SLIDE 1: TITLE SLIDE ====================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(8.0), Inches(2.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p0 = tf1.paragraphs[0]
    p0.text = "🛡️ NetGuard-IDS"
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TEAL
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf1.add_paragraph()
    p1.text = "Network Intrusion Detection & Automated Port Risk Analyzer"
    p1.font.size = Pt(18)
    p1.font.color.rgb = COLOR_WHITE
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "\nA Unified Dual-Engine Defensive Cybersecurity Platform"
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED
    p2.alignment = PP_ALIGN.CENTER

    meta_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(8.0), Inches(1.5))
    tf_meta = meta_box.text_frame
    p_meta = tf_meta.paragraphs[0]
    p_meta.text = "Author / Presenter: Aditya Yadav (BitWiseAditya27)\nDepartment of Computer Science & Engineering\nAcademic Year 2026"
    p_meta.font.size = Pt(12)
    p_meta.font.color.rgb = COLOR_MUTED
    p_meta.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 2: PROBLEM STATEMENT ====================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)
    add_header(s2, "The Cybersecurity Dilemma in Enterprise Networks")

    content_box = s2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2))
    tf2 = content_box.text_frame
    tf2.word_wrap = True

    points_s2 = [
        ("Explosion of Perimeter Vulnerabilities", "Organizations inadvertently expose dangerous services (SMB 445, Telnet 23, RDP 3389, unauthenticated Redis/Mongo) that serve as primary vectors for ransomware and lateral movement."),
        ("Limitations of Standalone Port Scanners (e.g. Nmap)", "Require administrative/root access for SYN scans, CLI-heavy, lack automated CVSS risk scoring, and provide zero intelligence on live traffic intrusions."),
        ("Fragility of Traditional Signature-based NIDS (e.g. Snort)", "Rely on rigid, static rules. High false positive rates, resource-heavy, and completely blind to zero-day flow anomalies."),
        ("The Solution -> NetGuard-IDS", "A unified, non-root system bridging socket-level port risk auditing with benchmarked machine learning traffic classification in one real-time dashboard.")
    ]
    for heading, desc in points_s2:
        p_h = tf2.add_paragraph()
        p_h.text = f"• {heading}: "
        p_h.font.bold = True
        p_h.font.size = Pt(14)
        p_h.font.color.rgb = COLOR_TEAL

        p_d = tf2.add_paragraph()
        p_d.text = f"   {desc}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 3: OBJECTIVES & SCOPE ====================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)
    add_header(s3, "Project Objectives & Core Deliverables")

    tf3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf3.word_wrap = True

    objs = [
        ("1. Multithreaded Non-Root Port Auditing", "Build asynchronous TCP connect scanner using Python socket + concurrent.futures with application banner grabbing."),
        ("2. Rule-Based CVSS Risk Engine", "Translate discovered ports & banners into risk tiers (Critical, High, Medium, Low) and calculate an aggregate 0–100 risk index."),
        ("3. Machine Learning Intrusion Classifier", "Train, compare, and deploy Random Forest & XGBoost pipelines on the NSL-KDD benchmark across 5 threat taxonomies."),
        ("4. Executive Defensive Dashboard", "Interactive Streamlit interface with Plotly gauges, attack distributions, SQLite audit history, and multi-format exports (PDF, JSON, CSV)."),
        ("5. Strict Ethical Guardrails", "Mandatory pre-scan consent, rate-limiting, and restriction to localhost, RFC1918 private subnets, and scanme.nmap.org.")
    ]
    for h, d in objs:
        p = tf3.add_paragraph()
        p.text = f"✔ {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_desc = tf3.add_paragraph()
        p_desc.text = f"   {d}\n"
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 4: SYSTEM ARCHITECTURE ====================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)
    add_header(s4, "Multi-Tier System Architecture")

    if os.path.exists(ARCH_IMG):
        s4.shapes.add_picture(ARCH_IMG, Inches(1.0), Inches(1.6), width=Inches(8.0))

    # ==================== SLIDE 5: MODULE 1 - PORT SCANNER & RISK ENGINE ====================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5)
    add_header(s5, "Defensive Engine 1: Port Scanner & Risk Engine")

    tf5 = s5.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf5.word_wrap = True

    m1_points = [
        ("4 Adaptive Scan Modes", "Quick (Top 100), Common (1-1024), Custom Range, and Full (1-65535) with rate limiting."),
        ("Application Banner Grabbing", "Safely queries open TCP sockets for daemon identification (HTTP, OpenSSH, FTP, MySQL, Redis, SMB)."),
        ("Heuristic Banner Escalation", "Automatically flags legacy/outdated services (e.g., SMBv1 EternalBlue flags, Anonymous FTP, old OpenSSH)."),
        ("CVSS Threat Modeling", "Associates open ports with known exploit vectors: WannaCry (445), BlueKeep (3389), Plaintext sniff (21/23), NoSQL wipes (27017)."),
        ("Overarching Verdict (0–100)", "Balances maximum CVSS severity with total threat exposure to assign LOW, MEDIUM, HIGH, or CRITICAL verdict.")
    ]
    for h, d in m1_points:
        p = tf5.add_paragraph()
        p.text = f"⚡ {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf5.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 6: MODULE 2 - MACHINE LEARNING IDS ====================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)
    add_header(s6, "Defensive Engine 2: ML Intrusion Detection Engine")

    tf6 = s6.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf6.word_wrap = True

    m2_points = [
        ("Benchmark Dataset", "NSL-KDD benchmark dataset with 41 flow features (connection duration, protocol, service, flags, bytes, host rates)."),
        ("Feature Engineering Pipeline", "StandardScaler for numerical flow dynamics + OneHotEncoder for protocol and service categories."),
        ("5 Attack Taxonomies Detected", "Normal (Benign), DoS (Neptune, Smurf), Probe (Nmap, Portsweep), R2L (Guess Password, FTP Write), U2R (Buffer Overflow, Rootkit)."),
        ("Dual-Algorithm Comparison", "Benchmarked Random Forest against XGBoost to determine optimal latency-accuracy trade-off.")
    ]
    for h, d in m2_points:
        p = tf6.add_paragraph()
        p.text = f"🧠 {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf6.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 7: ML EVALUATION & COMPARISON ====================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7)
    add_header(s7, "Model Comparison & Benchmark Results")

    table_shape = s7.shapes.add_table(3, 5, Inches(0.8), Inches(1.8), Inches(8.4), Inches(1.8))
    t = table_shape.table
    headers = ["Classifier", "Accuracy", "Weighted F1", "Precision", "Recall"]
    for i, h in enumerate(headers):
        cell = t.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_TEAL
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.color.rgb = COLOR_BG

    row1 = ["Random Forest (Selected)", "98.30%", "0.9826", "0.9838", "0.9830"]
    row2 = ["XGBoost Classifier", "98.25%", "0.9821", "0.9830", "0.9825"]
    for i, val in enumerate(row1):
        c = t.cell(1, i)
        c.text = val
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD
        c.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE
        c.text_frame.paragraphs[0].font.size = Pt(12)

    for i, val in enumerate(row2):
        c = t.cell(2, i)
        c.text = val
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD
        c.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE
        c.text_frame.paragraphs[0].font.size = Pt(12)

    tf7_extra = s7.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(8.4), Inches(2.8)).text_frame
    tf7_extra.word_wrap = True
    p_ex = tf7_extra.paragraphs[0]
    p_ex.text = "Key Observations & Class Breakdown:"
    p_ex.font.bold = True
    p_ex.font.size = Pt(14)
    p_ex.font.color.rgb = COLOR_TEAL

    observations = [
        "• Normal Traffic & R2L Attacks achieved 100% precision and recall.",
        "• DoS attacks achieved 95% precision and 100% recall (zero missed flooding attempts).",
        "• Sub-100ms inference time per batch makes it suitable for real-time flow evaluation.",
        "• Serialization with Joblib allows instantaneous model loading in serverless environments."
    ]
    for obs in observations:
        p_o = tf7_extra.add_paragraph()
        p_o.text = obs
        p_o.font.size = Pt(12)
        p_o.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 8: SAFETY & ETHICAL GOVERNANCE ====================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8)
    add_header(s8, "Safety Engineering & Ethical Compliance")

    tf8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf8.word_wrap = True

    eth = [
        ("Mandatory Pre-Scan Consent Checkbox", "User must explicitly affirm ownership or written testing authorization before the audit button triggers."),
        ("Programmatic Target Allowlisting", "Only localhost (127.0.0.1), private RFC1918 subnets (10.x, 172.16-31.x, 192.168.x), and scanme.nmap.org are permitted. Arbitrary public IPs are strictly blocked."),
        ("Hosted Demo Mode", "Generates authentic simulated scan results in cloud environments where raw socket outbound traffic is blocked by cloud provider firewalls."),
        ("Zero Command Injection Surface", "Built strictly with Python native sockets; no subprocess execution with shell=True."),
        ("Regulatory Compliance", "Designed in adherence with CFAA (US), Computer Misuse Act (UK), and IT Act 2000 (India).")
    ]
    for h, d in eth:
        p = tf8.add_paragraph()
        p.text = f"⚖️ {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf8.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 9: LIVE DEMO & REPOSITORIES ====================
    s9 = prs.slides.add_slide(blank_layout)
    apply_background(s9)
    add_header(s9, "Deployment, Testing & Live Demonstration")

    tf9 = s9.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf9.word_wrap = True

    demo_pts = [
        ("Public GitHub Repository", "https://github.com/BitWiseAditya27/netguard-ids"),
        ("Live Streamlit Cloud Deployment", "https://netguard-ids.streamlit.app/"),
        ("Automated Testing Suite", "10/10 Pytest tests passing covering validators, CVSS math, socket handlers, and ML inference."),
        ("Multi-Version CI/CD", "GitHub Actions workflow running linting and tests across Python 3.10, 3.11, and 3.12 on every push."),
        ("Export Capabilities", "Instant export of security findings to JSON, CSV, and Markdown executive summaries.")
    ]
    for h, d in demo_pts:
        p = tf9.add_paragraph()
        p.text = f"🌐 {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf9.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 10: CONCLUSION & FUTURE WORK ====================
    s10 = prs.slides.add_slide(blank_layout)
    apply_background(s10)
    add_header(s10, "Conclusion & Future Enhancements")

    tf10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf10.word_wrap = True

    conc = [
        ("Conclusion", "NetGuard-IDS provides a comprehensive, production-grade cybersecurity tool combining lightweight host scanning, actionable CVSS-style risk scoring, and 98.3% accurate ML flow classification in an intuitive dashboard."),
        ("Future Scope 1: Live Packet Sniffing", "Incorporate eBPF / Scapy for real-time promiscuous interface packet capture on hardware taps."),
        ("Future Scope 2: Deep Learning Autoencoders", "Deploy unsupervised LSTM Autoencoders to detect zero-day evasion techniques."),
        ("Future Scope 3: Enterprise SIEM Integration", "Integrate automated webhooks and alerts into Splunk, Elastic SIEM, and Discord/Slack SOC channels."),
        ("Thank You!", "Questions & Answers are welcome.")
    ]
    for h, d in conc:
        p = tf10.add_paragraph()
        p.text = f"📌 {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf10.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    prs.save(OUTPUT_PPTX)
    print(f"Presentation saved successfully to: {OUTPUT_PPTX}")

if __name__ == "__main__":
    build_presentation()
