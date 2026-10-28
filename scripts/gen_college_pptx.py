"""
NetGuard-IDS College Submission Presentation Generator.
Generates College_Presentation.pptx following the exact 13-slide academic template:
Title, Introduction, Problem Statement, Objectives, System Methodology/Workflow,
Operating Environment, Non-Functional Requirements, User Classes, Advantages,
Limitations, Future Scope, Conclusion, Thank You.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DOCS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
WF_IMG = os.path.join(DOCS_DIR, "workflow_diagram.png")
OUTPUT_PPTX = os.path.join(DOCS_DIR, "College_Presentation.pptx")

# Academic Dark Cyan Palette
COLOR_BG = RGBColor(14, 17, 23)        # #0E1117
COLOR_CARD = RGBColor(26, 29, 36)      # #1A1D24
COLOR_TEAL = RGBColor(0, 212, 178)     # #00D4B2
COLOR_WHITE = RGBColor(245, 245, 245)
COLOR_MUTED = RGBColor(170, 175, 185)

def apply_background(slide):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title, slide_num):
    # Header tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(7.0), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = "NETGUARD-IDS | DEFENSIVE CYBERSECURITY & ML"
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_TEAL

    # Slide number
    num_box = slide.shapes.add_textbox(Inches(8.5), Inches(0.4), Inches(0.8), Inches(0.35))
    tf_n = num_box.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.text = str(slide_num)
    p_n.font.size = Pt(12)
    p_n.font.bold = True
    p_n.font.color.rgb = COLOR_MUTED
    p_n.alignment = PP_ALIGN.RIGHT

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(8.4), Inches(0.7))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_WHITE

def build_college_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==================== SLIDE 1: TITLE SLIDE ====================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(8.0), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p0 = tf1.paragraphs[0]
    p0.text = "NetGuard-IDS"
    p0.font.size = Pt(40)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TEAL
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf1.add_paragraph()
    p1.text = "Network Intrusion Detection & Port Risk Analyzer"
    p1.font.size = Pt(18)
    p1.font.color.rgb = COLOR_WHITE
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "An AI-Driven Dual-Engine Defensive Cybersecurity Platform"
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED
    p2.alignment = PP_ALIGN.CENTER

    team_box = s1.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(8.0), Inches(2.2))
    tf_team = team_box.text_frame
    p_by = tf_team.paragraphs[0]
    p_by.text = "Presentation by:"
    p_by.font.size = Pt(13)
    p_by.font.bold = True
    p_by.font.color.rgb = COLOR_TEAL
    p_by.alignment = PP_ALIGN.CENTER

    members = [
        "Aditya Yadav",
        "Ashutosh Mehta",
        "Atishay Jain"
    ]
    for m in members:
        pm = tf_team.add_paragraph()
        pm.text = m
        pm.font.size = Pt(14)
        pm.font.color.rgb = COLOR_WHITE
        pm.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 2: INTRODUCTION ====================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)
    add_header(s2, "Introduction", 2)

    tf2 = s2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf2.word_wrap = True

    pts2 = [
        "Rapid expansion of digital infrastructure and remote networks has created widespread perimeter visibility gaps.",
        "Manual tracking of open network ports, services, and unauthorized access attempts is impractical at enterprise scale.",
        "Need for an automated, unified defensive system to support network engineers, security auditors, and SOC analysts.",
        "Proposed Solution: NetGuard-IDS — A Machine Learning & socket-based platform to discover exposed services, calculate automated CVSS risk verdicts, and detect network intrusions in real time."
    ]
    for p in pts2:
        par = tf2.add_paragraph()
        par.text = f"• {p}\n"
        par.font.size = Pt(14)
        par.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 3: PROBLEM STATEMENT ====================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)
    add_header(s3, "Problem Statement", 3)

    tf3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf3.word_wrap = True

    pts3 = [
        "Enterprise environments process millions of network flows daily, making manual threat triage impossible.",
        "Traditional scanners (e.g. Nmap) require root privileges, manual CLI flags, and lack built-in CVSS risk scoring.",
        "Classic signature-based NIDS (e.g. Snort) fail to identify novel attack variations and zero-day traffic anomalies.",
        "Lack of unified coordination between port vulnerability discovery and real-time flow intrusion detection.",
        "Results in critical missed opportunities to isolate exposed vectors (SMB 445, RDP 3389, Telnet 23) before ransomware compromise occurs."
    ]
    for p in pts3:
        par = tf3.add_paragraph()
        par.text = f"• {p}\n"
        par.font.size = Pt(14)
        par.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 4: OBJECTIVES ====================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)
    add_header(s4, "Objectives", 4)

    tf4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf4.word_wrap = True

    pts4 = [
        "Develop an automated, non-root multithreaded TCP port scanner with application banner grabbing.",
        "Construct an automated CVSS-style risk engine that scores vulnerabilities from 0–100 with concrete remediation steps.",
        "Preprocess flow telemetry and train benchmark ML models (Random Forest & XGBoost) on the NSL-KDD dataset.",
        "Detect and classify 5 attack taxonomies: Normal, DoS, Probe, R2L (Remote-to-Local), and U2R (User-to-Root).",
        "Provide real-time interactive telemetry through a modern, responsive Streamlit dashboard.",
        "Enforce strict ethical guardrails and target whitelisting (RFC 1918 & scanme.nmap.org) to prevent malicious use."
    ]
    for p in pts4:
        par = tf4.add_paragraph()
        par.text = f"• {p}\n"
        par.font.size = Pt(13)
        par.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 5: SYSTEM METHODOLOGY (WORKFLOW) ====================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5, )
    add_header(s5, "System Methodology & End-to-End Workflow", 5)

    if os.path.exists(WF_IMG):
        s5.shapes.add_picture(WF_IMG, Inches(0.8), Inches(1.5), width=Inches(8.4))

    # Add small bulleted summary below image
    tf5_sub = s5.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(8.4), Inches(1.4)).text_frame
    tf5_sub.word_wrap = True
    p5_sub = tf5_sub.paragraphs[0]
    p5_sub.text = "1. Data Collection & Target Ingestion -> 2. Safety Whitelist -> 3. Dual Engine (Port Scanner & ML Flow Classifier) -> 4. CVSS Scoring -> 5. Streamlit Dashboard & Report Export"
    p5_sub.font.size = Pt(11)
    p5_sub.font.bold = True
    p5_sub.font.color.rgb = COLOR_TEAL

    # ==================== SLIDE 6: OPERATING ENVIRONMENT ====================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)
    add_header(s6, "Operating Environment", 6)

    tf6 = s6.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf6.word_wrap = True

    pts6 = [
        ("Core Language & Runtime", "Python 3.10+ (Cross-platform compatibility across Windows, Linux, macOS)."),
        ("Data Processing & Machine Learning", "Pandas, NumPy, Scikit-learn, XGBoost, Joblib."),
        ("Networking & Concurrency", "Python standard socket library, concurrent.futures.ThreadPoolExecutor."),
        ("Presentation & Visualization", "Streamlit 1.35+, Plotly interactive charts."),
        ("Persistence & Database", "Embedded SQLite 3 (Thread-safe connection pool, zero configuration)."),
        ("Deployment & Cloud Hosting", "Streamlit Community Cloud (Hosted serverless persistent container)."),
        ("Browser Support", "Google Chrome, Microsoft Edge, Mozilla Firefox, Apple Safari.")
    ]
    for h, d in pts6:
        p = tf6.add_paragraph()
        p.text = f"• {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf6.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 7: NON-FUNCTIONAL REQUIREMENTS ====================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7)
    add_header(s7, "Non-Functional Requirements", 7)

    tf7 = s7.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf7.word_wrap = True

    pts7 = [
        ("High-Speed Performance", "Sub-second port discovery for top ports; sub-100ms ML batch inference latency."),
        ("Zero Root Privilege Dependency", "Operates entirely within user space via unprivileged TCP connect sockets."),
        ("Strict Target Authorization & Safety", "Programmatic allowlisting blocking arbitrary public IPs; mandatory consent checkbox."),
        ("High Reliability & Fault Tolerance", "Graceful timeout handling for firewalled/dropped packets without app crashes."),
        ("Multi-Format Exportability", "One-click generation of audit records in JSON, CSV, and Markdown formats."),
        ("Scalability", "Thread-pool architecture capable of scaling across subnets and cloud demo environments.")
    ]
    for h, d in pts7:
        p = tf7.add_paragraph()
        p.text = f"✔ {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf7.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 8: USER CLASSES ====================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8)
    add_header(s8, "User Classes", 8)

    tf8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf8.word_wrap = True

    pts8 = [
        ("Security Engineers & SOC Analysts", "Conduct host posture audits, upload live flow captures, inspect intrusion breakdowns, and monitor perimeter health."),
        ("System & Network Administrators", "Identify exposed internal services (e.g. unauthenticated Redis, SMBv1), review CVSS severity scores, and apply hardening advice."),
        ("Penetration Testers & Auditors", "Perform authorized vulnerability assessments with exportable compliance logs and evidence artifacts."),
        ("Academic Evaluators & Students", "Study defensive cybersecurity architectures, port risk mechanics, and ML-based flow threat detection.")
    ]
    for h, d in pts8:
        p = tf8.add_paragraph()
        p.text = f"👤 {h}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEAL
        p_d = tf8.add_paragraph()
        p_d.text = f"   {d}\n"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 9: ADVANTAGES ====================
    s9 = prs.slides.add_slide(blank_layout)
    apply_background(s9)
    add_header(s9, "Advantages", 9)

    tf9 = s9.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf9.word_wrap = True

    pts9 = [
        "Automated Identification: Seamlessly bridges port scanning and flow classification in one platform.",
        "Non-Root Portability: Unlike Nmap SYN scans, requires zero sudo or root administrative privileges.",
        "High Predictive Accuracy: 98.30% accuracy and 0.9826 F1-score achieved on NSL-KDD benchmark.",
        "Actionable Remediation: Explains exact real-world exploit vectors (WannaCry, BlueKeep) and hardening steps.",
        "Embedded Persistence: Zero-setup SQLite tracking for historical trend comparisons.",
        "Ethical Guardrails: Native controls prevent abuse against unauthorized external targets."
    ]
    for p in pts9:
        par = tf9.add_paragraph()
        par.text = f"• {p}\n"
        par.font.size = Pt(13)
        par.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 10: LIMITATIONS ====================
    s10 = prs.slides.add_slide(blank_layout)
    apply_background(s10)
    add_header(s10, "Limitations", 10)

    tf10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf10.word_wrap = True

    pts10 = [
        "Predictive Accuracy Depends on Flow Features: Model relies on 41 NSL-KDD attributes; unformatted raw PCAPs require conversion.",
        "Cloud Hosting Socket Restrictions: Outbound TCP connects may be blocked by serverless cloud providers (mitigated via Hosted Demo Mode).",
        "Passive vs. Active Interception: Currently performs detection and classification rather than inline packet dropping (IPS).",
        "Firewall SYN Drops: Scanning heavily firewalled hosts can increase duration if connect timeouts are set high."
    ]
    for p in pts10:
        par = tf10.add_paragraph()
        par.text = f"• {p}\n"
        par.font.size = Pt(13)
        par.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 11: FUTURE SCOPE ====================
    s11 = prs.slides.add_slide(blank_layout)
    apply_background(s11)
    add_header(s11, "Future Scope", 11)

    tf11 = s11.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf11.word_wrap = True

    pts11 = [
        "Live Packet Sniffing Integration: Implement eBPF or Scapy promiscuous capture directly from hardware interfaces.",
        "Deep Learning Zero-Day Detection: Incorporate unsupervised LSTM Autoencoders to detect uncataloged attack variants.",
        "Active Intrusion Prevention (IPS): Add automated iptables/Windows Firewall rule injection to drop attacking IPs in real time.",
        "Enterprise SIEM Webhooks: Export alerts directly into Splunk, Elastic SIEM, and SOC messaging channels (Slack/Discord)."
    ]
    for p in pts11:
        par = tf11.add_paragraph()
        par.text = f"• {p}\n"
        par.font.size = Pt(13)
        par.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 12: CONCLUSION ====================
    s12 = prs.slides.add_slide(blank_layout)
    apply_background(s12)
    add_header(s12, "Conclusion", 12)

    tf12 = s12.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2)).text_frame
    tf12.word_wrap = True

    pts12 = [
        "Proactive perimeter auditing and continuous traffic inspection are essential for modern cybersecurity defense.",
        "NetGuard-IDS successfully demonstrates that non-root socket auditing and machine learning can be unified into an efficient, lightweight solution.",
        "Achieves 98.30% threat classification accuracy with actionable CVSS remediation advice and interactive SOC telemetry.",
        "Provides educational and industrial organizations with an accessible, production-ready cyber defense environment."
    ]
    for p in pts12:
        par = tf12.add_paragraph()
        par.text = f"• {p}\n"
        par.font.size = Pt(14)
        par.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 13: THANK YOU ====================
    s13 = prs.slides.add_slide(blank_layout)
    apply_background(s13)

    t_box13 = s13.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(8.0), Inches(3.0))
    tf13 = t_box13.text_frame
    
    p_ty = tf13.paragraphs[0]
    p_ty.text = "Thank You!"
    p_ty.font.size = Pt(44)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_TEAL
    p_ty.alignment = PP_ALIGN.CENTER

    p_qa = tf13.add_paragraph()
    p_qa.text = "\nQuestions & Feedback are Welcome"
    p_qa.font.size = Pt(18)
    p_qa.font.color.rgb = COLOR_WHITE
    p_qa.alignment = PP_ALIGN.CENTER

    p_repo = tf13.add_paragraph()
    p_repo.text = "\nGitHub: github.com/BitWiseAditya27/netguard-ids\nLive Demo: netguard-ids.streamlit.app"
    p_repo.font.size = Pt(12)
    p_repo.font.color.rgb = COLOR_MUTED
    p_repo.alignment = PP_ALIGN.CENTER

    prs.save(OUTPUT_PPTX)
    print(f"College presentation saved to: {OUTPUT_PPTX}")

if __name__ == "__main__":
    build_college_presentation()
