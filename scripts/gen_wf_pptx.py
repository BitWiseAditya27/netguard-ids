"""
NetGuard-IDS Workflow Architecture Presentation Generator.
Dedicated multi-slide presentation detailing the complete system workflow step-by-step:
1. Title: End-to-End Workflow Architecture
2. Workflow High-Level Blueprint (Diagram)
3. Step 1: Target Ingestion & Consent Verification
4. Step 2: Ethical Whitelisting & Safety Policy
5. Step 3: Non-Root Multithreaded Port Scanner
6. Step 4: Rule-Based CVSS Scoring & Threat Modeling
7. Step 5: NSL-KDD Machine Learning Flow Classifier
8. Step 6: Telemetry Dashboard & Multi-Format Reporting
9. Thank You
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DOCS_DIR = r"C:\Users\py056\.gemini\antigravity\scratch\netguard-ids\docs"
OUTPUT_PPTX = os.path.join(DOCS_DIR, "Workflow_Architecture_Presentation.pptx")
WF_IMG = os.path.join(DOCS_DIR, "workflow_diagram.png")

COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)
COLOR_TITLE = RGBColor(15, 23, 42)
COLOR_ACCENT = RGBColor(30, 58, 138)        # #1E3A8A Navy
COLOR_MUTED = RGBColor(100, 116, 139)

def set_slide_background_white(slide):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_WHITE
    bg_shape.line.fill.background()
    return bg_shape

def add_workflow_slide(prs, title_text, slide_number, bullet_points, subheader=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_white(slide)

    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(0.9))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_ACCENT

    # Thin line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(8.4), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    line.line.fill.background()

    # Content Box
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(8.4), Inches(4.8))
    tf_c = content_box.text_frame
    tf_c.word_wrap = True

    if subheader:
        p_sub = tf_c.paragraphs[0]
        p_sub.text = subheader
        p_sub.font.size = Pt(14)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_MUTED
        p_sub.space_after = Pt(12)
        start_idx = 1
    else:
        start_idx = 0

    for i, pt in enumerate(bullet_points):
        p = tf_c.add_paragraph() if (start_idx == 1 or i > 0) else tf_c.paragraphs[0]
        p.text = pt
        p.font.size = Pt(15)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(14)

    # Number
    num_box = slide.shapes.add_textbox(Inches(8.5), Inches(6.8), Inches(0.8), Inches(0.4))
    p_n = num_box.text_frame.paragraphs[0]
    p_n.text = str(slide_number)
    p_n.font.size = Pt(12)
    p_n.font.color.rgb = COLOR_MUTED
    p_n.alignment = PP_ALIGN.RIGHT

    return slide

def generate_workflow_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_white(s1)

    by_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(4.2), Inches(3.5))
    tf_by = by_box.text_frame
    p_by = tf_by.paragraphs[0]
    p_by.text = "Presentation by-"
    p_by.font.size = Pt(18)
    p_by.font.bold = True
    p_by.font.color.rgb = COLOR_ACCENT
    p_by.space_after = Pt(10)
    for name in ["Aditya Yadav", "Ashutosh Mehta", "Atishay Jain"]:
        pn = tf_by.add_paragraph()
        pn.text = name
        pn.font.size = Pt(16)
        pn.font.color.rgb = COLOR_TEXT_MAIN
        pn.space_after = Pt(6)

    title_box = s1.shapes.add_textbox(Inches(5.0), Inches(2.0), Inches(4.3), Inches(3.5))
    tf_title = title_box.text_frame
    pt1 = tf_title.paragraphs[0]
    pt1.text = "NetGuard-IDS"
    pt1.font.size = Pt(32)
    pt1.font.bold = True
    pt1.font.color.rgb = COLOR_ACCENT
    pt1.space_after = Pt(8)
    pt2 = tf_title.add_paragraph()
    pt2.text = "Complete Architectural Workflow & Data Pipeline"
    pt2.font.size = Pt(18)
    pt2.font.bold = True
    pt2.font.color.rgb = COLOR_TITLE
    pt2.space_after = Pt(12)
    pt3 = tf_title.add_paragraph()
    pt3.text = "Deep-Dive into Ingestion, Ethical Gates, Dual Scanning Engines, and Report Synthesis"
    pt3.font.size = Pt(13)
    pt3.font.color.rgb = COLOR_MUTED

    # Slide 2: High-Level Architecture Diagram
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_white(s2)
    t_box2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.4), Inches(0.8))
    p2 = t_box2.text_frame.paragraphs[0]
    p2.text = "End-to-End Architectural Workflow Blueprint"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT
    line2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(8.4), Inches(0.04))
    line2.fill.solid()
    line2.fill.fore_color.rgb = COLOR_ACCENT
    line2.line.fill.background()

    if os.path.exists(WF_IMG):
        s2.shapes.add_picture(WF_IMG, Inches(0.8), Inches(1.6), width=Inches(8.4))

    num_box2 = s2.shapes.add_textbox(Inches(8.5), Inches(6.8), Inches(0.8), Inches(0.4))
    num_box2.text_frame.paragraphs[0].text = "2"
    num_box2.text_frame.paragraphs[0].font.color.rgb = COLOR_MUTED

    # Slide 3: Step 1 - Ingestion & Consent
    add_workflow_slide(
        prs,
        title_text="Workflow Step 1: Target Ingestion & Consent",
        slide_number=3,
        subheader="Stage 1: User Parameters & Mandatory Consent Policy",
        bullet_points=[
            "User supplies target identifier: IP Address (127.0.0.1, 192.168.1.1) or Hostname (scanme.nmap.org).",
            "Scan mode parameter configured: Quick (Top 100), Common (1-1024), Custom Range, or Full (1-65535).",
            "Mandatory Pre-Scan Verification Checkbox: Programmatically blocks scan initiation until ownership or written authorization is confirmed.",
            "Network flow CSV ingestion for intrusion detection module."
        ]
    )

    # Slide 4: Step 2 - Ethical Whitelisting
    add_workflow_slide(
        prs,
        title_text="Workflow Step 2: Ethical Whitelisting & Safety",
        slide_number=4,
        subheader="Stage 2: Programmatic Enforcement & Anti-Abuse Rules",
        bullet_points=[
            "Target resolution through socket.gethostbyname with regex sanitization to block command injection.",
            "RFC 1918 Private Subnet Verification: Verifies target belongs to 10.0.0.0/8, 172.16.0.0/12, or 192.168.0.0/16.",
            "Public IP Quarantine: Any unapproved public internet IP is immediately halted with a policy violation notice.",
            "Official research target exception: Specifically permits authorized host scanme.nmap.org.",
            "Rate-limiting and thread capping (10-100 max) to prevent denial of service."
        ]
    )

    # Slide 5: Step 3 - Non-Root Port Scanner Engine
    add_workflow_slide(
        prs,
        title_text="Workflow Step 3: Non-Root Multithreaded Scanner",
        slide_number=5,
        subheader="Stage 3A: Socket-Level Reconnaissance & Banner Grabbing",
        bullet_points=[
            "ThreadPoolExecutor dispatches asynchronous non-root socket.connect_ex probes.",
            "Operates without sudo/root or raw packet privileges across Windows and Linux.",
            "Active Banner Grabbing: Queries open sockets for software version strings (HTTP, SSH, ProFTPD, MySQL, SMB).",
            "Heuristic Banner Analysis: Detects obsolete services, anonymous FTP permits, and legacy SMBv1 headers.",
            "Cloud Fallback (Hosted Demo Mode): Simulates enterprise host telemetry when serverless outbound sockets are blocked."
        ]
    )

    # Slide 6: Step 4 - CVSS Risk Engine
    add_workflow_slide(
        prs,
        title_text="Workflow Step 4: Rule-Based CVSS Risk Engine",
        slide_number=6,
        subheader="Stage 4A: Vulnerability Mapping & Aggregated Verdicts",
        bullet_points=[
            "Maps discovered ports to CVSS v3.1 base metrics and severity categories (Critical, High, Medium, Low).",
            "Associates findings with known exploit vectors (WannaCry on 445, BlueKeep on 3389, cleartext sniff on 21/23).",
            "Mathematical Score Aggregation: Score = min(100, (raw_score * 0.8) + (max_cvss * 3.5)).",
            "Issues overarching security verdict: LOW, MEDIUM, HIGH, or CRITICAL.",
            "Generates concrete, actionable remediation steps for system administrators."
        ]
    )

    # Slide 7: Step 5 - ML Traffic Classifier Engine
    add_workflow_slide(
        prs,
        title_text="Workflow Step 5: Machine Learning Flow Classifier",
        slide_number=7,
        subheader="Stage 3B: Dual-Model Intrusion Detection Pipeline",
        bullet_points=[
            "Input: Network flow telemetry comprising 41 NSL-KDD attributes (duration, bytes, flags, rates).",
            "Preprocessing: StandardScaler for continuous dynamics + OneHotEncoder for categorical services/protocols.",
            "Inference Pipeline: Evaluates flows using trained Random Forest & XGBoost classifiers.",
            "5 Attack Taxonomies Classified: Normal, DoS (flooding), Probe (recon), R2L (unauthorized access), U2R (privilege escalation).",
            "Achieves 98.30% Accuracy and 0.9826 Weighted F1-Score in sub-100ms batch execution."
        ]
    )

    # Slide 8: Step 6 - SOC Telemetry & Reporting
    add_workflow_slide(
        prs,
        title_text="Workflow Step 6: Telemetry Dashboard & Reporting",
        slide_number=8,
        subheader="Stage 5: Visualization, Persistence & Export Artifacts",
        bullet_points=[
            "Streamlit Reactive Interface: Visualizes risk with Plotly gauge meters, severity donuts, and attack bar charts.",
            "Embedded Relational Persistence: Automatically stores scan sessions and ML metrics in SQLite 3.",
            "Multi-Format Artifact Export: Exports audit records to JSON, CSV, and Markdown executive summaries.",
            "Live SOC Feed: Real-time tracking of historical scans, malicious flow rates, and host threat status."
        ]
    )

    # Slide 9: Thank You
    s9 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_white(s9)
    t_box9 = s9.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(8.0), Inches(2.5))
    tf9 = t_box9.text_frame
    p_ty = tf9.paragraphs[0]
    p_ty.text = "Thank You!"
    p_ty.font.size = Pt(44)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_ACCENT
    p_ty.alignment = PP_ALIGN.CENTER
    p_qa = tf9.add_paragraph()
    p_qa.text = "\nNetGuard-IDS Workflow Architecture"
    p_qa.font.size = Pt(18)
    p_qa.font.color.rgb = COLOR_TITLE
    p_qa.alignment = PP_ALIGN.CENTER

    prs.save(OUTPUT_PPTX)
    print(f"Generated Workflow Architecture presentation at: {OUTPUT_PPTX}")

if __name__ == "__main__":
    generate_workflow_presentation()
